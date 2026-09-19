import io
from PIL import Image
import docx
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas


def create_synthetic_pdf(num_pages: int = 3) -> bytes:
    """Generates a valid multi-page PDF with semantically distinct topics on each page."""
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)

    page_contents = [
        (
            "NEXUS Processor Specification - Oryon CPU",
            "The 12-core Qualcomm Oryon CPU delivers high performance desktop multitasking.",
            "Each CPU cluster operates at 3.8 GHz peak frequency with multi-threaded efficiency.",
        ),
        (
            "NEXUS AI Accelerator - Hexagon NPU",
            "The dedicated Qualcomm Hexagon NPU delivers 45 TOPS of continuous AI matrix computation.",
            "This tensor engine executes quantized INT8 and INT4 neural networks without draining battery.",
        ),
        (
            "NEXUS Graphics System - Adreno GPU",
            "The Qualcomm Adreno GPU delivers 4.6 TFLOPS of graphics processing and OpenCL compute.",
            "High resolution multi-monitor display output is driven by the integrated display processor.",
        ),
    ]

    for p in range(num_pages):
        content = page_contents[p % len(page_contents)]
        c.drawString(100, 750, f"Page {p + 1}: {content[0]}")
        c.drawString(100, 720, content[1])
        c.drawString(100, 690, content[2])
        c.showPage()

    c.save()
    buffer.seek(0)
    return buffer.getvalue()


def create_scanned_synthetic_pdf() -> bytes:
    """Generates a valid PDF with an empty canvas (no digital text streams) simulating a scanned page."""
    buffer = io.BytesIO()
    c = canvas.Canvas(buffer, pagesize=letter)
    # Draw only a non-text vector rectangle
    c.rect(50, 50, 500, 700)
    c.showPage()
    c.save()
    buffer.seek(0)
    return buffer.getvalue()


def create_synthetic_docx() -> bytes:
    """Generates a valid DOCX file with headings, paragraphs, and a structured table."""
    doc = docx.Document()
    doc.add_heading("NEXUS Work Agent Architecture", level=1)
    doc.add_paragraph("NEXUS is an offline-first multimodal AI work agent.")
    doc.add_paragraph("All documents are processed locally without internet transmission.")

    # Add a table
    table = doc.add_table(rows=3, cols=2)
    table.cell(0, 0).text = "Modality"
    table.cell(0, 1).text = "Engine"
    table.cell(1, 0).text = "Speech"
    table.cell(1, 1).text = "Whisper-Small"
    table.cell(2, 0).text = "Vision"
    table.cell(2, 1).text = "CLIP ViT-B/32"

    buffer = io.BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer.getvalue()


def create_synthetic_txt() -> bytes:
    """Generates a valid UTF-8 plain text file with multiple paragraphs."""
    content = (
        "NEXUS Offline Local Ingestion Subsystem.\n\n"
        "Paragraph 1: Files are parsed and validated strictly on the local PC.\n"
        "No network calls or telemetry are allowed.\n\n"
        "Paragraph 2: Extracted chunks retain page numbers, filenames, and character counts.\n"
        "This enables deterministic source attribution."
    )
    return content.encode("utf-8")


def create_synthetic_png(width: int = 250, height: int = 250) -> bytes:
    """Generates a valid RGB PNG image."""
    img = Image.new("RGB", (width, height), color=(0, 229, 255))
    buffer = io.BytesIO()
    img.save(buffer, format="PNG")
    buffer.seek(0)
    return buffer.getvalue()


def create_synthetic_jpg(width: int = 200, height: int = 200) -> bytes:
    """Generates a valid JPEG image."""
    img = Image.new("RGB", (width, height), color=(255, 128, 0))
    buffer = io.BytesIO()
    img.save(buffer, format="JPEG")
    buffer.seek(0)
    return buffer.getvalue()


def create_invalid_file() -> bytes:
    """Generates invalid binary content that does not conform to any allowed magic signatures."""
    return b"\x00\x11\x22\x33\x44\x55\x66\x77\x88\x99\xaa\xbb\xcc\xdd\xee\xff"
